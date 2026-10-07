"""Inspect or export the question-only RH workbook without modifying it."""

import argparse
import csv
import json
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


DEFAULT_WORKBOOK = Path(__file__).resolve().parents[1] / "RH_dataset.xlsx"


def load_questions(path=DEFAULT_WORKBOOK):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook["Sheet2"]
        rows = worksheet.iter_rows(values_only=True)
        header = next(rows)
        positions = {name: header.index(name) for name in ("no", "category", "question")}
        questions = []
        for row in rows:
            question = row[positions["question"]]
            if not isinstance(question, str) or not question.strip():
                continue
            questions.append({
                "no": row[positions["no"]],
                "category": row[positions["category"]],
                "question": question,
            })
        return questions
    finally:
        workbook.close()


def export_questions(questions, path):
    path = Path(path)
    if path.suffix.lower() not in (".csv", ".jsonl"):
        raise ValueError("Output must end with .csv or .jsonl.")
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        if path.suffix.lower() == ".csv":
            writer = csv.DictWriter(handle, fieldnames=("no", "category", "question"))
            writer.writeheader()
            writer.writerows(questions)
        else:
            for question in questions:
                handle.write(json.dumps(question, ensure_ascii=False) + "\n")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--category", type=int, choices=(1, 2))
    parser.add_argument("--head", type=int, default=3)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    if args.head < 0:
        parser.error("--head must be nonnegative.")
    questions = load_questions(args.workbook)
    if args.category is not None:
        questions = [q for q in questions if q["category"] == args.category]
    print(f"Questions: {len(questions)}")
    print(f"Category counts: {dict(Counter(q['category'] for q in questions))}")
    for question in questions[:args.head]:
        print(json.dumps(question, ensure_ascii=False))
    if args.output:
        export_questions(questions, args.output)
        print(f"Exported to {args.output}")


if __name__ == "__main__":
    main()
